"""Offline smoke verification of the frozen app and its spawned workers."""
import json
import multiprocessing
import queue
import tempfile
from dataclasses import asdict
from pathlib import Path


def run(report_path=None):
    report = {"success": False, "checks": []}
    try:
        import pymupdf as fitz
        from docx import Document
        from .jobs import Task, worker
        from .storage import Credentials
        Credentials.backend()  # verify Windows backend is packaged, without saving keys
        report["checks"].append("credential_backend_import")
        # Explicit report directory also hosts scratch data in restricted build environments.
        scratch_parent = Path(report_path).resolve().parent if report_path else None
        with tempfile.TemporaryDirectory(prefix="DocumentStudio-selftest-", dir=scratch_parent) as folder:
            root = Path(folder)
            source = root / "source.pdf"
            with fitz.open() as doc:
                page = doc.new_page()
                page.insert_text((30, 40), "Document Studio smoke test")
                doc.save(source)
            word = root / "source.docx"
            document = Document()
            document.add_paragraph("Hello world")
            document.save(word)
            ctx = multiprocessing.get_context("spawn")
            for operation, input_file in (("extract_text", source), ("pdf_to_word", source), ("word_format", word)):
                output = root / operation
                output.mkdir()
                events, cancel = ctx.Queue(), ctx.Event()
                task = Task("selftest", operation, [str(input_file)], str(output), {"find": "Hello", "replace": "Goodbye"})
                process = ctx.Process(target=worker, args=(asdict(task), events, cancel, str(root / "cache"), str(output / ".stage")))
                process.start()
                done = None
                try:
                    # Drain during execution, preventing large IPC result deadlocks.
                    import time
                    deadline = time.monotonic() + 40
                    while time.monotonic() < deadline:
                        try:
                            message = events.get(timeout=.2)
                            if message[0] == "done":
                                done = message
                                break
                            if message[0] in ("failed", "cancelled"):
                                raise RuntimeError(f"{operation}: {message[1]}")
                        except queue.Empty:
                            if not process.is_alive():
                                break
                    process.join(timeout=5)
                    if done is None or process.exitcode != 0:
                        raise RuntimeError(f"{operation}: child process failed")
                    if not done[1] or not Path(done[1][0]).is_file():
                        raise RuntimeError(f"{operation}: missing output")
                    if operation == "word_format":
                        if Document(done[1][0]).paragraphs[0].text != "Goodbye world":
                            raise RuntimeError("word_format: output text mismatch")
                    report["checks"].append(operation + "_spawn_output")
                finally:
                    if process.is_alive():
                        process.terminate()
                        process.join()
                    process.close()
                    events.close()
        report["success"] = True
    except Exception as exc:
        report["error"] = str(exc)
    if report_path:
        Path(report_path).write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return 0 if report["success"] else 1
