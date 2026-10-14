import json
import zipfile
from threading import Event
from pathlib import Path
import numpy as np
import pymupdf as fitz
import httpx
import pytest
from docx import Document
from document_studio.pdf_tools import process_pdf, render_processed, pdf_to_word
from document_studio.api import APIClient, TranslationValidationError
from document_studio.translation import translate_docx


def test_rotated_image_protection_and_search_layer(tmp_path):
    source = tmp_path / "rotated.pdf"
    pix = fitz.Pixmap(fitz.csRGB, 20, 20, bytes([220, 230, 210]) * 400, False)
    with fitz.open() as doc:
        page = doc.new_page(width=100, height=120)
        page.draw_rect(page.rect, fill=(.9, .9, .85), color=None)
        page.insert_text((10, 60), "Search word", fontsize=8)
        page.insert_image(fitz.Rect(10, 10, 30, 30), pixmap=pix)
        page.set_rotation(90)
        doc.save(source)
    with fitz.open(source) as doc:
        image_rect = fitz.Rect(doc[0].get_image_info()[0]["bbox"]) * doc[0].rotation_matrix
        processed = render_processed(doc[0], 72, preserve_images=True)
        array = np.frombuffer(processed.samples, dtype=np.uint8).reshape(processed.height, processed.width, 3)
        assert (array[int(image_rect.y0 + 5), int(image_rect.x0 + 5)] == [220, 230, 210]).all()
    stage = tmp_path / "stage"
    stage.mkdir()
    output = process_pdf("whiten", [source], stage, {"dpi": 72, "keep_text": True}, lambda *a: None, Event())[0]
    with fitz.open(output) as doc:
        assert doc[0].rotation == 0
        assert len(doc[0].search_for("Search word")) == 1
        assert doc[0].get_text().count("Search word") == 1
        rendered = doc[0].get_pixmap(colorspace=fitz.csRGB, alpha=False)
        pixels = np.frombuffer(rendered.samples, dtype=np.uint8).reshape(rendered.height, rendered.width, 3)
        assert (pixels[0, 0] == 255).all()


def test_numeric_omission_rejected_and_fallback_batch(tmp_path):
    client = APIClient({"provider": "openai", "base_url": "https://example.com/v1"}, "dummy",
        transport=httpx.MockTransport(lambda r: httpx.Response(200, json={"choices": [{"message": {"content": '{"items":[{"id":"p0","text":"没有数字"}]}'}}]})))
    client.pause = lambda *a: None
    with pytest.raises(TranslationValidationError):
        client.translate("m", [{"id": "p0", "text": "Value is 123.45"}], "English", "Chinese", "")
    client.close()
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("first paragraph")
    doc.add_paragraph("second paragraph")
    doc.save(source)
    class Translator:
        base = "https://example.com/v1"
        provider = "openai"
        calls = []
        def translate(self, model, entries, *args):
            self.calls.append(len(entries))
            if len(entries) > 1:
                raise TranslationValidationError("injected batch failure")
            return [{"id": entries[0]["id"], "text": "译：" + entries[0]["text"]}]
    translator = Translator()
    translate_docx(source, tmp_path / "done.docx", {"model": "m", "source_lang": "English", "target_lang": "Chinese"}, translator,
                   tmp_path / "cache", lambda *a: None, Event())
    assert translator.calls == [2, 1, 1]
    assert len(json.loads(next((tmp_path / "cache").glob("*.json")).read_text(encoding="utf-8"))["translations"]) == 2


def test_pdf_conversion_uses_strict_page_errors(tmp_path, monkeypatch):
    source = tmp_path / "source.pdf"
    with fitz.open() as doc:
        page = doc.new_page()
        page.insert_text((10, 30), "Native PDF text")
        doc.save(source)
    import pdf2docx
    class Converter:
        def __init__(self, *a):
            pass
        def convert(self, *a, **options):
            assert options["ignore_page_error"] is False
            assert options["raw_exceptions"] is True
            raise RuntimeError("injected parsing failure")
        def close(self):
            pass
    monkeypatch.setattr(pdf2docx, "Converter", Converter)
    with pytest.raises(RuntimeError):
        pdf_to_word(source, tmp_path / "result.docx", lambda *a: None, Event())
