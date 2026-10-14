from dataclasses import asdict
from queue import Queue, Empty
from threading import Event
from pathlib import Path
from docx import Document
import pymupdf as fitz
from document_studio.jobs import Task, worker


def messages(events):
    result = []
    while True:
        try:
            result.append(events.get_nowait())
        except Empty:
            return result


def test_pdf_export_failure_preserves_paid_translation_and_checkpoint(tmp_path, monkeypatch):
    source = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("Hello world")
    doc.save(source)
    output = tmp_path / "out"
    output.mkdir()
    class Client:
        base = "https://example.com/v1"
        provider = "openai"
        def __init__(self, *args, **kwargs):
            pass
        def translate(self, model, entries, *args):
            return [{"id": row["id"], "text": "译：" + row["text"]} for row in entries]
        def close(self):
            pass
    monkeypatch.setattr("document_studio.api.APIClient", Client)
    monkeypatch.setattr("document_studio.jobs.Credentials.get", lambda *a: "dummy")
    monkeypatch.setattr("document_studio.word_tools.detect_engines", lambda *a: {"word": True, "libreoffice": None})
    def fail(*a):
        raise ValueError("injected export failure")
    monkeypatch.setattr("document_studio.word_tools.word_to_pdf", fail)
    options = {"profile": {"id": "synthetic"}, "model": "m", "source_lang": "English", "target_lang": "Chinese", "output_pdf": True}
    events = Queue()
    worker(asdict(Task("translation", "translate", [str(source)], str(output), options)), events, Event(), str(tmp_path / "cache"), str(output / ".stage"))
    records = messages(events)
    assert any(row[0] == "failed" and "DOCX 已保存" in row[1] for row in records)
    assert any(row[0] == "output" for row in records)
    assert "译：" in Document(next(output.glob("*.docx"))).paragraphs[0].text
    assert list((tmp_path / "cache").glob("*.json"))


def test_pdf_preflight_stops_before_credentials_or_api(tmp_path, monkeypatch):
    monkeypatch.setattr("document_studio.word_tools.detect_engines", lambda *a: {"word": False, "libreoffice": None})
    def forbid(*a):
        raise AssertionError("Credential access must not occur")
    monkeypatch.setattr("document_studio.jobs.Credentials.get", forbid)
    events = Queue()
    task = Task("translation", "translate", [], str(tmp_path), {"profile": {"id": "dummy"}, "output_pdf": True})
    worker(asdict(task), events, Event(), str(tmp_path / "cache"), str(tmp_path / ".stage"))
    records = messages(events)
    assert any(row[0] == "failed" and "尚未调用翻译 API" in row[1] for row in records)
