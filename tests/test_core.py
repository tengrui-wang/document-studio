import json
import multiprocessing
import zipfile
from pathlib import Path
from threading import Event

import httpx
import numpy as np
import pymupdf as fitz
import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

from document_studio.api import APIClient, APIError, validate_base_url
from document_studio.common import Cancelled, page_numbers, publish
from document_studio.jobs import Task, worker
from document_studio.pdf_tools import pdf_to_word, process_pdf, render_processed
from document_studio.storage import Settings
from document_studio.translation import translate_docx
from document_studio.word_tools import word_format


def progress(*args):
    pass


@pytest.fixture
def pdf(tmp_path):
    path = tmp_path / "source.pdf"
    with fitz.open() as doc:
        for text in ("First page", "Second page", "Third page"):
            page = doc.new_page(width=300, height=400)
            page.draw_rect(page.rect, color=None, fill=(.88, .94, .84))
            page.insert_text((30, 60), text)
        doc.save(path)
    return path


def test_pages_and_publish_no_overwrite(tmp_path):
    assert page_numbers("3,1-2", 3) == [2, 0, 1]
    for invalid in ("0", "1,1", "4", "3-1", "1-", "1,,2"):
        with pytest.raises(ValueError):
            page_numbers(invalid, 3)
    original = tmp_path / "original.docx"
    original.write_bytes(b"source")
    temp = tmp_path / "stage"
    temp.write_bytes(b"output")
    output = publish(temp, tmp_path, "original", ".docx")
    assert output.name == "original_1.docx"
    assert original.read_bytes() == b"source"
    assert output.read_bytes() == b"output"


def test_pdf_operations(pdf, tmp_path):
    stage = tmp_path / "stage"
    stage.mkdir()
    outputs = process_pdf("extract_pages", [pdf], stage, {"pages": "3,1"}, progress, Event())
    with fitz.open(outputs[0]) as doc:
        assert len(doc) == 2
        assert "Third page" in doc[0].get_text()
    outputs = process_pdf("rotate", [pdf], stage, {"pages": "2", "angle": 90}, progress, Event())
    with fitz.open(outputs[0]) as doc:
        assert [page.rotation for page in doc] == [0, 90, 0]
    parts = process_pdf("split", [pdf], stage, {"split_size": 2}, progress, Event())
    assert len(parts) == 2
    merged = process_pdf("merge", parts, stage, {}, progress, Event())
    with fitz.open(merged[0]) as doc:
        assert len(doc) == 3
    text = process_pdf("extract_text", [pdf], stage, {}, progress, Event())[0].read_text(encoding="utf-8")
    assert all(word in text for word in ("First page", "Second page", "Third page"))
    event = Event()
    event.set()
    with pytest.raises(Cancelled):
        process_pdf("merge", [pdf], stage, {}, progress, event)


def test_whitening(pdf, tmp_path):
    with fitz.open(pdf) as doc:
        pix = render_processed(doc[0], 72, 180, preserve_images=False)
        pixels = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, 3)
        assert (pixels[5, 5] == 255).all()
        assert pixels.min() < 80  # dark text survives
    stage = tmp_path / "stage"
    stage.mkdir()
    output = process_pdf("whiten", [pdf], stage, {"pages": "1", "dpi": 72}, progress, Event())[0]
    with fitz.open(output) as doc:
        assert len(doc) == 1
        assert len(doc[0].get_images()) == 1


def test_image_preservation_and_extraction(tmp_path):
    pix = fitz.Pixmap(fitz.csRGB, 20, 20, bytes([220, 230, 210]) * 400, False)
    path = tmp_path / "images.pdf"
    with fitz.open() as doc:
        page = doc.new_page(width=100, height=100)
        page.insert_image(fitz.Rect(10, 10, 30, 30), pixmap=pix)
        doc.save(path)
    with fitz.open(path) as doc:
        kept = render_processed(doc[0], 72, 180, preserve_images=True)
        erased = render_processed(doc[0], 72, 180, preserve_images=False)
        a = np.frombuffer(kept.samples, dtype=np.uint8).reshape(100, 100, 3)
        b = np.frombuffer(erased.samples, dtype=np.uint8).reshape(100, 100, 3)
        assert (a[20, 20] == [220, 230, 210]).all()
        assert (b[20, 20] == 255).all()
    stage = tmp_path / "stage"
    stage.mkdir()
    assert process_pdf("extract_images", [path], stage, {}, progress, Event())[0].stat().st_size > 0
    with pytest.raises(ValueError, match="扫描"):
        pdf_to_word(path, tmp_path / "bad.docx", progress, Event())


def make_word(tmp_path):
    source = tmp_path / "sample.docx"
    doc = Document()
    doc.add_heading("Title", level=1)
    p = doc.add_paragraph()
    p.add_run("Hello ").bold = True
    p.add_run("world")
    formula = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    text = OxmlElement("m:t")
    text.text = "x=2"
    run.append(text)
    formula.append(run)
    p._p.append(formula)
    p.add_run(" after equation")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Cell"
    doc.sections[0].header.paragraphs[0].text = "Header"
    pix = fitz.Pixmap(fitz.csRGB, 10, 10, bytes([255, 0, 0]) * 100, False)
    png = tmp_path / "image.png"
    pix.save(png)
    doc.add_picture(str(png), width=Inches(.2))
    doc.save(source)
    return source


def test_word_format_keeps_structure(tmp_path):
    source = make_word(tmp_path)
    dest = tmp_path / "formatted.docx"
    word_format(source, dest, {"font": "Arial", "font_size": 12, "find": "Hello world", "replace": "Goodbye world", "heading_size": 18}, progress, Event())
    doc = Document(dest)
    assert "Goodbye world" in doc.paragraphs[1].text
    assert len(doc.tables) == 1 and len(doc.inline_shapes) == 1
    assert len(list(doc.element.iter(qn("m:oMath")))) == 1
    assert doc.paragraphs[0].runs[0].font.size.pt == 18
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(dest) as changed:
        image = next(name for name in original.namelist() if name.startswith("word/media/"))
        assert original.read(image) == changed.read(image)


class FakeTranslator:
    base = "https://example.com/v1"
    provider = "openai"
    def __init__(self, fail_after=None):
        self.calls = 0
        self.fail_after = fail_after
    def translate(self, model, entries, source, target, glossary):
        self.calls += 1
        if self.fail_after is not None and self.calls > self.fail_after:
            raise APIError("simulated failure")
        return [{"id": entry["id"], "text": "译：" + entry["text"]} for entry in entries]


def test_translation_resume_and_protected_content(tmp_path):
    source = make_word(tmp_path)
    dest = tmp_path / "translated.docx"
    options = {"model": "test", "source_lang": "English", "target_lang": "Chinese", "chunk_size": 200}
    client = FakeTranslator()
    translate_docx(source, dest, options, client, tmp_path / "cache", progress, Event())
    assert client.calls > 0
    second = FakeTranslator(fail_after=0)
    translate_docx(source, tmp_path / "resumed.docx", options, second, tmp_path / "cache", progress, Event())
    assert second.calls == 0
    doc = Document(dest)
    assert "译：" in doc.paragraphs[1].text
    assert "译：" in doc.tables[0].cell(0, 0).text
    assert "译：" in doc.sections[0].header.paragraphs[0].text
    assert len(doc.inline_shapes) == 1
    assert next(doc.element.iter(qn("m:t"))).text == "x=2"
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(dest) as changed:
        for name in original.namelist():
            if name.startswith("word/media/"):
                assert original.read(name) == changed.read(name)


def test_resume_after_failure(tmp_path):
    source = tmp_path / "long.docx"
    doc = Document()
    for index in range(5):
        doc.add_paragraph(f"paragraph-{index} " * 10)
    doc.save(source)
    options = {"model": "test", "source_lang": "English", "target_lang": "Chinese", "chunk_size": 200}
    with pytest.raises(APIError):
        translate_docx(source, tmp_path / "failed.docx", options, FakeTranslator(1), tmp_path / "cache", progress, Event())
    assert not (tmp_path / "failed.docx").exists()
    client = FakeTranslator()
    translate_docx(source, tmp_path / "done.docx", options, client, tmp_path / "cache", progress, Event())
    assert client.calls == 4
    changed = dict(options, model="different")
    client2 = FakeTranslator()
    translate_docx(source, tmp_path / "other.docx", changed, client2, tmp_path / "cache", progress, Event())
    assert client2.calls == 5


def test_api_validation_sanitization_and_usage():
    profile = {"provider": "openai", "base_url": "https://example.com/v1"}
    usages = []
    def handle(request):
        assert request.headers["authorization"] == "Bearer TEST_SECRET"
        if request.url.path.endswith("models"):
            return httpx.Response(200, json={"data": [{"id": "model-a"}]})
        payload = json.loads(request.content)
        assert payload["model"] == "model-a"
        return httpx.Response(200, json={"choices": [{"message": {"content": '{"items":[{"id":"p0","text":"你好"}]}'}, "finish_reason": "stop"}], "usage": {"prompt_tokens": 8, "completion_tokens": 3}})
    client = APIClient(profile, "TEST_SECRET", transport=httpx.MockTransport(handle), usage=lambda a, b: usages.append((a, b)))
    assert client.models() == ["model-a"]
    assert client.translate("model-a", [{"id": "p0", "text": "hello"}], "English", "Chinese", "")[0]["text"] == "你好"
    assert usages == [(8, 3)]
    client.close()
    error_client = APIClient(profile, "TEST_SECRET", transport=httpx.MockTransport(lambda r: httpx.Response(401, json={"error": "TEST_SECRET"})))
    with pytest.raises(APIError) as captured:
        error_client.models()
    assert "TEST_SECRET" not in str(captured.value)
    error_client.close()
    for bad in ("http://example.com", "https://key@example.com", "https://example.com?key=secret", "https://example.com/#secret"):
        with pytest.raises(ValueError):
            validate_base_url(bad)


def test_api_rejects_wrong_order_and_redirect():
    profile = {"provider": "openai", "base_url": "https://example.com/v1"}
    client = APIClient(profile, "test", transport=httpx.MockTransport(lambda r: httpx.Response(200, json={"choices": [{"message": {"content": '{"items":[{"id":"wrong","text":"text"}]}'}}]})))
    client.pause = lambda seconds: None
    with pytest.raises(APIError, match="校验"):
        client.translate("m", [{"id": "p0", "text": "hello"}], "English", "Chinese", "")
    client.close()
    redirected = APIClient(profile, "test", transport=httpx.MockTransport(lambda r: httpx.Response(302, headers={"location": "https://untrusted.example"})))
    with pytest.raises(APIError, match="302"):
        redirected.models()
    redirected.close()


def test_settings_excludes_keys(tmp_path):
    settings = Settings(tmp_path)
    settings.data["profiles"] = [{"id": "1", "name": "demo", "provider": "openai", "base_url": "https://example.com/v1", "models": [], "api_key": "TEST_SECRET"}]
    settings.save()
    assert "TEST_SECRET" not in settings.path.read_text()


def test_real_pdf_to_word(pdf, tmp_path):
    target = tmp_path / "converted.docx"
    pdf_to_word(pdf, target, progress, Event())
    doc = Document(target)
    assert "First page" in " ".join(node.text or "" for node in doc.element.iter(qn("w:t")))


def test_spawn_worker(pdf, tmp_path):
    ctx = multiprocessing.get_context("spawn")
    events, cancel = ctx.Queue(), ctx.Event()
    out = tmp_path / "out"
    out.mkdir()
    task = Task("extract", "extract_text", [str(pdf)], str(out))
    from dataclasses import asdict
    process = ctx.Process(target=worker, args=(asdict(task), events, cancel, str(tmp_path / "cache"), str(out / ".stage")))
    process.start()
    process.join(timeout=30)
    if process.is_alive():
        process.terminate()
        process.join()
        pytest.fail("worker timed out")
    assert process.exitcode == 0
    messages = []
    while True:
        try:
            messages.append(events.get(timeout=.2))
        except Exception:
            break
    assert any(item[0] == "done" for item in messages), messages
    assert next(out.glob("*.txt")).exists()
    assert not (out / ".stage").exists()
    events.close()
    process.close()
